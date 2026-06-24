#!/usr/bin/env python3
"""Generates the FlexForge deliverable documents (DOCX + XLSX). PDFs are produced from
these via LibreOffice (soffice --headless --convert-to pdf). Run from the flexforge/ dir:
    python3 docs/generate_docs.py
Outputs into docs/generated/.
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

OUT = os.path.join(os.path.dirname(__file__), "generated")
os.makedirs(OUT, exist_ok=True)

ACCENT = RGBColor(0x4A, 0x6B, 0xFF)

# ---------------------------------------------------------------------------
# 1) TECHNICAL DOCUMENTATION (DOCX)
# ---------------------------------------------------------------------------

def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    return h

def add_para(doc, text, bold=False, italic=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    return p

def add_bullets(doc, items):
    for it in items:
        doc.add_paragraph(str(it), style="List Bullet")

def add_numbered(doc, items):
    for it in items:
        doc.add_paragraph(str(it), style="List Number")

def add_table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
    return t

def code_block(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Courier New"
    run.font.size = Pt(9)

def build_tech_doc():
    doc = Document()

    title = doc.add_heading("FlexForge — Technical Documentation", 0)
    sub = add_para(doc, "Runtime-interpretation no-code application platform")
    sub.runs[0].italic = True
    add_para(doc, "Version 0.1.0  ·  Generated 2026-06-24  ·  Audience: engineering, QA, onboarding")

    add_heading(doc, "1. Executive Overview", 1)
    add_para(doc,
        "FlexForge is a packaged runtime-interpretation engine: one hardened Spring Boot "
        "image reads an immutable config bundle (entities, fields, auth, flows) and BECOMES "
        "a full backend application at runtime. No per-app source code is generated. The same "
        "metadata is projected onto multiple protocols (REST, GraphQL, SSE, WebSocket), and "
        "all logic is expressed as no-code flows.")
    add_para(doc, "Key properties:", bold=True)
    add_bullets(doc, [
        "No-code for the builder: apps are config + flows, never hand-written code.",
        "One interpreted engine, one config bundle (the Mendix model) — not multi-stack codegen.",
        "Modular monolith, one engine process per app.",
        "AI-ready: AI modules slot in as flow steps (uniform FlowStep interface).",
        "Multi-language via escape hatches (serverless connector, JS code-step, WASM) — not multiple engines.",
    ])

    add_heading(doc, "2. High-Level Design (HLD)", 1)
    add_para(doc, "Two planes:")
    add_bullets(doc, [
        "Design Platform (future) — authors config bundles (visual builder, AI assist).",
        "Runtime Engine (this repo) — interprets a bundle and serves the app.",
    ])
    add_para(doc, "Request lifecycle (runtime):")
    code_block(doc,
        "HTTP (REST / GraphQL / SSE / WebSocket / Flow trigger)\n"
        "        |\n"
        "        v\n"
        "  Protocol surface  (DynamicRestController / GraphQLController / Realtime / FlowController)\n"
        "        |\n"
        "        v\n"
        "  AccessGuard (auth)  ->  DynamicCrudService  <->  MetadataRegistry (in-memory model)\n"
        "        |                        |\n"
        "        |                        v\n"
        "        |                   SqlDialect -> parameterized SQL -> HikariCP -> Database\n"
        "        v\n"
        "  EntityEvent stream -> SSE / WebSocket subscribers")
    add_para(doc, "Startup sequence:")
    add_numbered(doc, [
        "Load config bundle (YAML/JSON), from classpath or FLEXFORGE_CONFIG_PATH.",
        "Validate (JSON Schema + semantic rules: one PK/entity, no dup fields...).",
        "Choose SQL dialect from config; build HikariCP datasource.",
        "Sync schema (create tables; additive column add; destructive changes refused).",
        "Expose enabled protocol surfaces; register flow steps; engine is live.",
    ])

    add_heading(doc, "3. Low-Level Design (LLD) — Modules", 1)
    add_table(doc,
        ["Module / Package", "Responsibility"],
        [
            ["config.model", "POJOs: AppConfig, EntityConfig, FieldConfig, SecurityConfig, FlowConfig, StepConfig, EndpointConfig"],
            ["config", "ConfigLoader (YAML/JSON), ConfigValidator (JSON Schema + semantic), MetadataRegistry"],
            ["schema", "SqlDialect (H2/Postgres/MySQL), Naming, SchemaManager (DDL sync)"],
            ["data", "DynamicCrudService (generic CRUD), QueryOptions/Filter, FieldValidator, Page"],
            ["crypto", "EncryptionService (AES-256-GCM, SHA-256, base64)"],
            ["auth", "JwtService, AuthContext, AccessGuard, SecurityFilter, AuthController"],
            ["flow + flow.steps", "FlowEngine, FlowContext, Expressions, FlowStep + StepRegistry; steps: set/condition/db/http/log/respond"],
            ["realtime", "EntityEvent, RealtimeService (SSE), WebSocketEventHandler/Config"],
            ["api.rest / api.graphql", "DynamicRestController, GraphQLSchemaBuilder + GraphQLController"],
            ["api.meta / api.file / api.endpoint / api.flow / api.realtime", "Meta+OpenAPI, file upload/download, custom endpoints, flow trigger, realtime endpoints"],
            ["bootstrap", "EngineConfiguration (composition root), EngineBootstrap, WebConfig (CORS)"],
        ])

    add_heading(doc, "4. Data Model & Field Types", 1)
    add_para(doc, "Entities and fields are declared in config. Supported field types:")
    add_table(doc,
        ["Type", "Stored as", "Notes"],
        [
            ["STRING / TEXT", "VARCHAR / TEXT", "length, validation (minLength/maxLength/pattern/email)"],
            ["INT / LONG", "INTEGER / BIGINT", "numeric validation (min/max)"],
            ["DOUBLE / DECIMAL", "DOUBLE / DECIMAL", ""],
            ["BOOLEAN", "BOOLEAN/TINYINT", ""],
            ["DATE / TIMESTAMP", "DATE / TIMESTAMP", "auto createdAt/updatedAt by convention"],
            ["JSON", "JSONB/JSON/text", "native per dialect"],
            ["REFERENCE", "BIGINT", "relation to another entity; app-level integrity"],
            ["FILE", "CLOB/TEXT", "upload/download/base64 via file API"],
            ["(any) encrypted:true", "ciphertext", "AES-256-GCM at rest, transparent on read"],
        ])

    add_heading(doc, "5. Application Flow (No-Code Flows)", 1)
    add_para(doc,
        "Flows are the no-code logic layer: a trigger plus a graph of typed steps over a "
        "shared context. Steps read prior results via ${input.x}, ${steps.id.field}, "
        "${vars.x} expressions. The uniform FlowStep interface is the extension point — AI "
        "modules and connectors plug in here with no engine changes.")
    add_para(doc, "Example flow (config):", bold=True)
    code_block(doc,
        "flows:\n"
        "  - name: grade\n"
        "    steps:\n"
        "      - { id: check, type: condition,\n"
        "          params: {left: \"${input.score}\", op: \">=\", right: 50},\n"
        "          then: pass, else: fail }\n"
        "      - { id: pass, type: respond, params: { body: { result: \"pass\" } } }\n"
        "      - { id: fail, type: respond, params: { body: { result: \"fail\" } } }")
    add_para(doc, "Trigger: POST /flows/grade/run {\"score\":70}  ->  {\"result\":\"pass\"}")
    add_para(doc, "Step types: set, condition, db (CRUD), http (outbound), log, respond.")

    add_heading(doc, "6. Protocols", 1)
    add_table(doc,
        ["Protocol", "Status", "Endpoint"],
        [
            ["REST", "Done", "/api/{entity}"],
            ["GraphQL", "Done", "/graphql (+ /graphql/schema)"],
            ["Server-Sent Events", "Done", "/realtime/stream"],
            ["WebSocket", "Done", "/ws/events"],
            ["Webhooks / custom endpoints", "Done", "/run/{name}, flow http step"],
            ["Flow trigger", "Done", "/flows/{name}/run"],
            ["GraphQL subscriptions / MCP", "Roadmap P1", "—"],
            ["gRPC / OData", "Roadmap P2", "—"],
            ["SOAP / GraphQL Federation", "Roadmap P3", "—"],
        ])

    add_heading(doc, "7. Security", 1)
    add_bullets(doc, [
        "Auth: JWT (HS256) via /auth/login, or API-key via X-API-Key header.",
        "Authorization: per-entity read/write role rules enforced by one AccessGuard across ALL protocols.",
        "Field-level encryption (AES-256-GCM) via encrypted:true.",
        "CORS enabled for API surfaces; secrets via ${ENV} placeholders (12-factor).",
        "Roadmap: row-level security, SSO/SAML/OIDC, audit log, multi-tenancy.",
    ])

    add_heading(doc, "8. How to Run", 1)
    add_para(doc, "Prerequisites: JDK 21, Maven 3.9+. From the flexforge/ directory:")
    code_block(doc,
        "# run on in-memory H2 (zero setup)\n"
        "mvn spring-boot:run\n"
        "# open the admin UI\n"
        "open http://localhost:8080/\n\n"
        "# point at Postgres/MySQL (same image)\n"
        "export DB_ENGINE=postgres DB_URL=jdbc:postgresql://localhost:5432/app \\\n"
        "       DB_USERNAME=app DB_PASSWORD=secret\n"
        "mvn spring-boot:run\n\n"
        "# run a different app with the same engine\n"
        "export FLEXFORGE_CONFIG_PATH=/path/to/app.yaml && mvn spring-boot:run\n\n"
        "# build a jar / OCI image (no Dockerfile, no source in image)\n"
        "mvn package\n"
        "mvn compile jib:build -Dimage=ghcr.io/you/flexforge-engine:0.1.0")
    add_para(doc, "Quick smoke test:")
    code_block(doc,
        "curl -s -X POST localhost:8080/api/Customer -H 'Content-Type: application/json' \\\n"
        "  -d '{\"email\":\"a@b.com\",\"fullName\":\"Ada\"}'\n"
        "curl -s localhost:8080/api/Customer\n"
        "curl -s -X POST localhost:8080/graphql -H 'Content-Type: application/json' \\\n"
        "  -d '{\"query\":\"{ customerList { id email } }\"}'\n"
        "curl -s localhost:8080/realtime/recent")

    add_heading(doc, "9. Testing", 1)
    add_para(doc,
        "Automated tests are JUnit 5 + Spring Boot @SpringBootTest integration tests that "
        "boot the full engine on H2 and drive the real HTTP surfaces. Run: mvn test. "
        "See the companion workbook 'FlexForge_Test_Cases_and_Report.xlsx' for the full "
        "case list and latest results (26 tests passing).")

    add_heading(doc, "10. Roadmap (summary)", 1)
    add_bullets(doc, [
        "P0.5: select/multiselect + computed fields (formula/lookup/rollup), N:N relations.",
        "P1: AI steps (LLM/RAG/classify/extract) + provider gateway + vector store; MCP protocol; serverless + JS code-step.",
        "P1.5: audit engine, row/field-level security, notifications, scheduler trigger.",
        "P2: visual builder, kanban/calendar/charts, connectors, multi-tenancy, gRPC/OData, i18n, environments.",
        "P3: WASM plugins, vertical modules (MLS, EBR, payroll), SOAP, federation, offline.",
    ])

    path = os.path.join(OUT, "FlexForge_Technical_Documentation.docx")
    doc.save(path)
    return path


# ---------------------------------------------------------------------------
# 2) TEST CASES + TEST REPORT (XLSX)
# ---------------------------------------------------------------------------

TESTS = [
    # id, suite, title, steps, expected, type, status
    ("TC-01", "EngineIntegrationTest", "REST CRUD round-trip",
     "POST/GET/LIST/PUT/DELETE /api/Customer", "201 create; reads back; 404 after delete", "Integration", "PASS"),
    ("TC-02", "EngineIntegrationTest", "GraphQL mutation + query",
     "createProduct mutation then productList query", "Record created and listed via GraphQL", "Integration", "PASS"),
    ("TC-03", "SecurityIntegrationTest", "Unauthenticated rejected",
     "GET /api/Note without token (security on)", "HTTP 401", "Security", "PASS"),
    ("TC-04", "SecurityIntegrationTest", "Role-gated write",
     "ADMIN can POST; VIEWER POST denied", "ADMIN 201; VIEWER 403; VIEWER read 200", "Security", "PASS"),
    ("TC-05", "SecurityIntegrationTest", "Bad credentials rejected",
     "POST /auth/login wrong password", "HTTP 401", "Security", "PASS"),
    ("TC-06", "FeaturesIntegrationTest", "Validation: bad email / missing required",
     "POST Customer invalid email; missing email", "HTTP 400 with errors[]", "Validation", "PASS"),
    ("TC-07", "FeaturesIntegrationTest", "Search operators",
     "GET /api/Customer?tier_like=gol and ?tier=silver", "Filtered results returned", "Functional", "PASS"),
    ("TC-08", "FeaturesIntegrationTest", "Metadata hides secrets",
     "GET /__meta", "appId + entities present; no jwtSecret/password", "Security", "PASS"),
    ("TC-09", "FeaturesIntegrationTest", "OpenAPI generated",
     "GET /__meta/openapi.json", "openapi=3.0.3 with paths", "Functional", "PASS"),
    ("TC-10", "FeaturesIntegrationTest", "Admin UI served",
     "GET / and /app.js", "200; contains 'FlexForge Admin'", "Functional", "PASS"),
    ("TC-11", "RelationsIntegrationTest", "Relation + auto timestamps",
     "Create Customer, create Order referencing it", "Order links; createdAt/updatedAt set", "Functional", "PASS"),
    ("TC-12", "RelationsIntegrationTest", "Referential integrity",
     "Create Order with missing customerId", "HTTP 400 with errors[]", "Validation", "PASS"),
    ("TC-13", "EndpointsIntegrationTest", "Missing required param",
     "GET /run/ping (no 'who')", "HTTP 400; error mentions 'who'", "Functional", "PASS"),
    ("TC-14", "EndpointsIntegrationTest", "Valid endpoint json action",
     "GET /run/ping?who=ada", "200; message=pong", "Functional", "PASS"),
    ("TC-15", "EndpointsIntegrationTest", "Missing required header",
     "POST /run/notify without X-API-Key", "HTTP 400; error mentions X-API-Key", "Functional", "PASS"),
    ("TC-16", "EndpointsIntegrationTest", "Wrong method",
     "GET /run/notify (expects POST)", "HTTP 405", "Functional", "PASS"),
    ("TC-17", "EndpointsIntegrationTest", "Header satisfies check",
     "POST /run/notify with X-API-Key", "200; accepted=true", "Functional", "PASS"),
    ("TC-18", "CryptoFileIntegrationTest", "Encrypted field round-trip",
     "Create Customer with ssn; read back", "ssn returned decrypted (plaintext)", "Security", "PASS"),
    ("TC-19", "CryptoFileIntegrationTest", "File upload/download/base64",
     "Upload to Document.attachment; download + ?format=base64", "Bytes match original", "Functional", "PASS"),
    ("TC-20", "EncryptionServiceTest", "AES round-trip + ciphertext differs",
     "encrypt/decrypt a value", "decrypt==plaintext; cipher!=plaintext", "Unit", "PASS"),
    ("TC-21", "EncryptionServiceTest", "Non-deterministic ciphertext",
     "encrypt same value twice", "different ciphertexts; both decrypt", "Unit", "PASS"),
    ("TC-22", "EncryptionServiceTest", "Base64 round-trip",
     "encode/decode bytes", "bytes preserved", "Unit", "PASS"),
    ("TC-23", "FlowIntegrationTest", "Flow expression interpolation",
     "POST /flows/greet/run {name:Ada}", "{message:'Hello Ada'}", "Functional", "PASS"),
    ("TC-24", "FlowIntegrationTest", "Flow db step creates record",
     "POST /flows/signup/run", "Record created; visible via REST", "Functional", "PASS"),
    ("TC-25", "FlowIntegrationTest", "Flow condition branching",
     "POST /flows/grade/run score=70 / 30", "pass / fail respectively", "Functional", "PASS"),
    ("TC-26", "RealtimeIntegrationTest", "CRUD emits change events",
     "Create Product; GET /realtime/recent", "create event present for Product", "Functional", "PASS"),
]

def thin_border():
    s = Side(style="thin", color="D0D0D0")
    return Border(left=s, right=s, top=s, bottom=s)

def build_test_workbook():
    wb = Workbook()

    # Sheet 1: Test Cases
    ws = wb.active
    ws.title = "Test Cases"
    headers = ["ID", "Suite", "Title", "Steps", "Expected Result", "Type", "Status"]
    ws.append(headers)
    head_fill = PatternFill("solid", fgColor="4A6BFF")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.border = thin_border()
    for t in TESTS:
        ws.append(list(t))
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = thin_border()
        status = row[6]
        status.font = Font(bold=True, color="1B8A4B" if status.value == "PASS" else "C0392B")
    widths = [8, 26, 30, 40, 42, 13, 9]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = "A2"

    # Sheet 2: Test Report (summary)
    rep = wb.create_sheet("Test Report")
    total = len(TESTS)
    passed = sum(1 for t in TESTS if t[6] == "PASS")
    failed = total - passed
    by_type = {}
    by_suite = {}
    for t in TESTS:
        by_type[t[5]] = by_type.get(t[5], 0) + 1
        by_suite[t[1]] = by_suite.get(t[1], 0) + 1

    rep.append(["FlexForge Test Report"])
    rep["A1"].font = Font(bold=True, size=16, color="4A6BFF")
    rep.append(["Generated", "2026-06-24"])
    rep.append(["Engine version", "0.1.0"])
    rep.append(["Framework", "JUnit 5 + Spring Boot @SpringBootTest (H2)"])
    rep.append([])
    rep.append(["Metric", "Value"])
    rep.append(["Total tests", total])
    rep.append(["Passed", passed])
    rep.append(["Failed", failed])
    rep.append(["Pass rate", f"{round(passed/total*100)}%"])
    rep.append([])
    rep.append(["By type", "Count"])
    for k, v in sorted(by_type.items()):
        rep.append([k, v])
    rep.append([])
    rep.append(["By suite", "Count"])
    for k, v in sorted(by_suite.items()):
        rep.append([k, v])

    for row in rep.iter_rows():
        for c in row:
            if c.value in ("Metric", "By type", "By suite", "Value", "Count"):
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = head_fill
    rep.column_dimensions["A"].width = 42
    rep.column_dimensions["B"].width = 18

    path = os.path.join(OUT, "FlexForge_Test_Cases_and_Report.xlsx")
    wb.save(path)
    return path


if __name__ == "__main__":
    p1 = build_tech_doc()
    p2 = build_test_workbook()
    print("Wrote:", p1)
    print("Wrote:", p2)
